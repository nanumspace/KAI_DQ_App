# -*- coding: utf-8 -*-
"""
GitHub Pages '자료실'(docs/resources.html)이 쓰는 파일을 기본 명세 판(v3)에서 만든다.

    python docs/build_resources.py            # → docs/resources/ (Pages 배포 때 워크플로가 실행, 저장소에는 넣지 않음)

만드는 것
  data.json                       자료실 화면이 읽는 데이터: 코호트·표·컬럼·코드표·규칙(코호트별 적용 여부)·견본 기대 결과
  rules-v3.csv                    검증 규칙 목록
  column-spec-v3.xlsx             컬럼 명세서 (모든 표)
  template-v3-<COHORT>.xlsx       엑셀 입력 양식 (안내 + 시트당 테이블 + 명세) — 앱의 '엑셀 양식'과 같은 모양
  template-v3-<COHORT>-csv.zip    CSV 입력 양식 묶음 (헤더만 + README.txt) — 앱의 'CSV 양식 묶음'과 같은 모양

모두 spec/ · rules/ · dq/reports/ 에서 생성하므로 명세나 규칙을 고치면 다음 배포에 저절로 따라온다.
의뢰사 원본 자료(raw/, spec/client_v1/)는 읽지도 싣지도 않는다.
"""
import csv, io, json, os, sys, zipfile
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill
import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs/resources")
EDITION = "v3"


def load(rel):
    return yaml.safe_load(open(os.path.join(ROOT, rel), encoding="utf-8"))


SPEC = load("spec/v3/kai_cdm_spec_v3.yaml")
TABLES = SPEC["tables"]
PROFILES = load("spec/v3/cohort_profiles_v3.yaml")["cohorts"]
CODELISTS = load("spec/v3/codelists_v3.yaml")
RULES_ST = load("rules/rules_structural_v3.yaml")["rules"]
RULES_SE = load("rules/rules_semantic_v3.yaml")["rules"]

SPEC_HEADER = ["테이블", "테이블(한글)", "순서", "컬럼", "컬럼(한글)", "타입", "필수", "키", "참조", "코드표", "범위", "패턴", "설명"]
GUIDE = ["K-AI 코호트 데이터 입력 양식 (엑셀)", "", "시트 하나가 테이블 하나입니다. 시트 이름은 바꾸지 마세요.",
         "첫 줄은 컬럼명입니다. 이름은 바꾸지 마세요.", "빈 값은 빈 칸으로 두고, 날짜는 날짜 셀 또는 YYYY-MM-DD 문자열로 넣습니다.",
         "컬럼 설명은 '명세' 시트를 보세요."]
CSV_README = "\n".join(["K-AI 코호트 데이터 입력 양식 (CSV)", "",
                        "- 파일 하나가 테이블 하나입니다. 파일 이름은 바꾸지 마세요.",
                        "- 첫 줄은 컬럼명입니다. 순서는 바꿔도 되지만 이름은 바꾸지 마세요.",
                        "- 인코딩 UTF-8, 빈 값은 빈 칸, 날짜는 YYYY-MM-DD, 타임스탬프는 YYYY-MM-DD HH:MM:SS 입니다.",
                        "- 컬럼 설명은 컬럼 명세서(xlsx)를 참고하세요.", ""])


def spec_rows(tables):
    """app/src/main/input.ts 의 specRows 와 같은 열 구성"""
    rows = [SPEC_HEADER]
    for t in tables:
        st = TABLES[t]
        for i, f in enumerate(st["fields"], 1):
            rng = f"{f['range'][0]} ~ {f['range'][1]}" if f.get("range") else ""
            rows.append([t, st["kor"], i, f["name"], f.get("kor", ""), f["type"], "Y" if f.get("required") else "",
                         f.get("key", ""), f.get("ref", ""), f.get("codelist", ""), rng, f.get("pattern", ""), f.get("desc", "")])
    return rows


def add_sheet(wb, title, rows, header=True, widths=None):
    ws = wb.create_sheet(title[:31])
    for r in rows:
        ws.append(r)
    if header and rows:
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8EEF6")
        ws.freeze_panes = "A2"
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    return ws


SPEC_WIDTHS = {"A": 24, "B": 14, "C": 6, "D": 30, "E": 14, "F": 9, "G": 6, "H": 12, "I": 34, "J": 32, "K": 10, "L": 18, "M": 80}


def cohort_rules(cohort):
    """이 코호트에서 실행되는 규칙 id (엔진의 v1·v3 경로와 같은 판정)"""
    tabs = set(PROFILES[cohort]["tables"])
    ids = [r["id"] for r in RULES_ST if r["table"] in tabs and (r.get("scope", "all") == "all" or cohort in r["scope"])]
    ids += [r["id"] for r in RULES_SE if r.get("scope", "all") == "all" or cohort in r["scope"]]
    return ids


def expected(kind, cohort):
    p = os.path.join(ROOT, f"dq/reports/{kind}_{EDITION}_{cohort}/summary.json")
    if not os.path.exists(p):
        return None
    s = json.load(open(p, encoding="utf-8"))
    return dict(rules=s["rules_total"], violations=s["violations_total"], failed=s["rules_failed"], skipped=s["rules_skipped"])


def main():
    os.makedirs(OUT, exist_ok=True)
    files = OrderedDict()

    # 컬럼 명세서 (모든 표)
    all_tables = list(TABLES)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    add_sheet(wb, "컬럼 명세", spec_rows(all_tables), widths=SPEC_WIDTHS)
    wb.save(os.path.join(OUT, "column-spec-v3.xlsx"))
    files["column_spec"] = "column-spec-v3.xlsx"

    # 코호트별 입력 양식
    cohorts = []
    for c, p in PROFILES.items():
        tabs = p["tables"]
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        add_sheet(wb, "안내", [[g] for g in GUIDE], header=False, widths={"A": 90})
        for t in tabs:
            add_sheet(wb, t, [[f["name"] for f in TABLES[t]["fields"]]])
        add_sheet(wb, "명세", spec_rows(tabs), widths=SPEC_WIDTHS)
        xlsx = f"template-v3-{c}.xlsx"
        wb.save(os.path.join(OUT, xlsx))
        zname = f"template-v3-{c}-csv.zip"
        with zipfile.ZipFile(os.path.join(OUT, zname), "w", zipfile.ZIP_DEFLATED) as z:
            for t in tabs:
                z.writestr(f"{c}/{t}.csv", "\ufeff" + ",".join(f["name"] for f in TABLES[t]["fields"]) + "\n")
            z.writestr(f"{c}/README.txt", CSV_README)
        rules = cohort_rules(c)
        exp = dict(clean=expected("clean", c), dirty=expected("dirty", c))
        if exp["clean"] and exp["clean"]["rules"] != len(rules):
            sys.exit(f"{c}: 규칙 수 판정이 엔진과 다르다 (자료실 {len(rules)} / 엔진 {exp['clean']['rules']})")
        cohorts.append(OrderedDict(id=c, kor=p["kor"], tables=tabs, rule_ids=rules, template_xlsx=xlsx, template_csv=zname,
                                   expected=exp))

    # 규칙 목록
    rules = []
    for r in RULES_ST:
        rules.append(OrderedDict(id=r["id"], kind="구조", name=r["name"], category=r["category"], subcategory=r.get("subcategory", ""),
                                 severity=r["severity"], table=r["table"], fields=",".join(r.get("fields") or []),
                                 check=r["check"], desc=r.get("desc", ""), sql=""))
    for r in RULES_SE:
        rules.append(OrderedDict(id=r["id"], kind="교차", name=r["name"], category=r["category"], subcategory=r.get("subcategory", ""),
                                 severity=r["severity"], table=",".join(r.get("requires") or []), fields="", check="sql",
                                 desc=r.get("desc", ""), sql=r["sql"].strip()))
    by_rule = {r["id"]: [c["id"] for c in cohorts if r["id"] in c["rule_ids"]] for r in rules}
    for r in rules:
        r["cohorts"] = by_rule[r["id"]]
    with open(os.path.join(OUT, "rules-v3.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        cols = ["id", "kind", "name", "category", "subcategory", "severity", "table", "fields", "check", "desc", "cohorts", "sql"]
        w.writerow(cols)
        for r in rules:
            w.writerow([";".join(r[k]) if k == "cohorts" else r[k] for k in cols])
    files["rules_csv"] = "rules-v3.csv"

    tables = OrderedDict()
    for t, st in TABLES.items():
        tables[t] = OrderedDict(kor=st["kor"], category=st["category"], pk=st["pk"], fields=[
            OrderedDict((k, f.get(k, "")) for k in ("name", "type", "required", "key", "ref", "codelist", "range", "pattern", "desc"))
            for f in st["fields"]])
    codelists = OrderedDict()
    for sect in CODELISTS.values():
        for k, v in sect.items():
            codelists[k] = OrderedDict((str(a), str(b)) for a, b in v.items())
    meta = SPEC["meta"]
    data = OrderedDict(edition=EDITION, spec_name=meta["name"], spec_version=meta["version"],
                       counts=dict(tables=len(TABLES), fields=sum(len(t["fields"]) for t in TABLES.values()),
                                   structural=len(RULES_ST), semantic=len(RULES_SE)),
                       files=files, cohorts=cohorts, tables=tables, codelists=codelists, rules=rules)
    for c in data["cohorts"]:
        del c["rule_ids"]
        c["rules"] = sum(1 for r in rules if c["id"] in r["cohorts"])
    json.dump(data, open(os.path.join(OUT, "data.json"), "w", encoding="utf-8"), ensure_ascii=False)
    print(f"자료실: 표 {data['counts']['tables']} · 필드 {data['counts']['fields']} · 규칙 {len(rules)} · 코호트 {len(cohorts)} → {os.path.relpath(OUT, ROOT)}")


if __name__ == "__main__":
    main()

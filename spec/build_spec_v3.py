# -*- coding: utf-8 -*-
"""
v3 명세: 의뢰사 9월 30일 매뉴얼에 2026-10-01 검토 결정(채택 14건)을 반영한 수정안을 기계 판독 명세로 만든다.

    python spec/build_spec_v3.py

입력  : spec/client_v1/매뉴얼_질환별 데이터 테이블_20260930_수정안.xlsx
        (빨간 취소선 행 = 삭제 결정이므로 명세에서 뺀다. 첫 시트 '수정 이력'은 읽지 않는다.)
출력  : spec/v3/kai_cdm_spec_v3.yaml   테이블·필드 (v1 명세와 같은 스키마 — 엔진의 v1 경로를 그대로 쓴다)
        spec/v3/codelists_v3.yaml      설명의 'N=라벨' 코드표 + v1 의 OMOP 집합(같은 표·필드에서 빌림)
        spec/v3/cohort_profiles_v3.yaml 코호트별 테이블 구성

v1 명세(spec/kai_cdm_spec.yaml)에 같은 표·같은 이름의 필드가 있으면 범위(range)·OMOP 코드표·외부 키(FK_EXTERNAL)를
빌린다. 수정안이 정한 형식(ICD-O-3 코드, AJCC 판수)은 아래 EXTRA 에 적는다.
"""
import io, os, re, sys
from collections import OrderedDict
import openpyxl
import yaml

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "spec/client_v1/매뉴얼_질환별 데이터 테이블_20260930_수정안.xlsx")
OUT = os.path.join(ROOT, "spec/v3")

CORE = ["PERSON", "VISIT_OCCURRENCE", "PROCEDURE_OCCURRENCE", "DRUG_EXPOSURE", "CONDITION_OCCURRENCE",
        "MEASUREMENT", "OBSERVATION", "NOTE"]
DISEASE = {"LUNG_CANCER": "폐암", "BREAST_CANCER": "유방암", "COLORECTAL_CANCER": "대장암",
           "MASLD": "대사이상지방간", "DIABETES": "당뇨병", "LYMPHOMA": "림프종"}
SACT_COHORTS = ["LUNG_CANCER", "BREAST_CANCER", "COLORECTAL_CANCER", "LYMPHOMA"]
EXTENSION = ["SACT", "ADVERSE_EVENT"]
REFERENCE = ["DRUG_INGREDIENT_STRUCTURE"]

TYPE = {"integer": "integer", "float": "float", "date": "date", "varchar": "varchar", "RBDMS dependent text": "text"}
# 매뉴얼은 FK 표시만 있고 대상이 없다. 이름으로 정한다.
FK_REF = {"person_id": "PERSON.person_id", "visit_occurrence_id": "VISIT_OCCURRENCE.visit_occurrence_id",
          "sact_id": "SACT.sact_id", "note_id": "NOTE.note_id"}
# 수정안이 정한 형식·범위 (검토 결정 C04·C06·C19, SACT 설명의 ECOG 0–5)
EXTRA = {
    ("LUNG_CANCER", "histology_icdo3_code"): dict(pattern=r"^\d{4}/[0-9]$"),
    ("LUNG_CANCER", "primary_site_icdo3_code"): dict(pattern=r"^C\d{2}\.\d$"),
    ("LYMPHOMA", "histology_icdo3_code"): dict(pattern=r"^\d{4}/[0-9]$"),
    ("LUNG_CANCER", "clinical_stage_ajcc_edition"): dict(codelist="LUNG_CANCER.clinical_stage_ajcc_edition"),
    ("SACT", "ecog_performance_status"): dict(range=[0, 5]),
    ("SACT", "line_of_therapy"): dict(range=[1, 20]),
    ("SACT", "number_of_cycles"): dict(range=[1, 100]),
}
EXTRA_CODELISTS = {"LUNG_CANCER.clinical_stage_ajcc_edition": {"6": "AJCC 6판", "7": "AJCC 7판", "8": "AJCC 8판"}}
CODE_LINE = re.compile(r"^\s*\d+\s*=")
CODE_ITEM = re.compile(r"(\d+)\s*=\s*([^,\n]+)")


def table_of(title):
    m = re.search(r"\(([A-Z_]+)\)\s*$", title)
    return m.group(1) if m else None


def category(t):
    return "core" if t in CORE else "disease" if t in DISEASE else "extension" if t in EXTENSION else "reference"


def main():
    v1 = yaml.safe_load(open(os.path.join(ROOT, "spec/kai_cdm_spec.yaml"), encoding="utf-8"))["tables"]
    v1_cl = yaml.safe_load(open(os.path.join(ROOT, "spec/codelists.yaml"), encoding="utf-8"))
    v1_omop = v1_cl.get("omop", {})
    wb = openpyxl.load_workbook(SRC)
    tables, embedded, omop, dropped = OrderedDict(), OrderedDict(), OrderedDict(), []
    for ws in wb.worksheets:
        t = table_of(ws.title)
        if t is None:
            continue
        section = str(ws["A1"].value or "").strip()
        fields, pk = [], None
        for r in range(3, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if name is None:
                continue
            name = str(name).strip()
            if ws.cell(r, 1).font.strike:
                dropped.append(f"{t}.{name}")
                continue
            typ = str(ws.cell(r, 2).value or "").strip()
            if typ not in TYPE:
                raise SystemExit(f"{ws.title} {r}행 {name}: 알 수 없는 타입 {typ!r}")
            key = str(ws.cell(r, 3).value or "").strip()
            req = str(ws.cell(r, 4).value or "").strip()
            if req not in ("Y", "N"):
                raise SystemExit(f"{ws.title} {r}행 {name}: 필수 여부가 Y/N 이 아님 ({req!r})")
            desc = str(ws.cell(r, 5).value or "").strip()
            f = OrderedDict(name=name, type=TYPE[typ], required=req == "Y", desc=desc)
            if key == "PK":
                f["key"], pk = "PK", name
            elif key == "FK":
                if name not in FK_REF:
                    raise SystemExit(f"{t}.{name}: FK 대상을 모름")
                f["key"], f["ref"] = "FK", FK_REF[name]
            base = next((x for x in v1.get(t, {}).get("fields", []) if x["name"] == name), None)
            if base:
                if not key and base.get("key") == "FK_EXTERNAL":
                    f["key"], f["ref"] = "FK_EXTERNAL", base["ref"]
                if base.get("range"):
                    f["range"] = list(base["range"])
                cl = base.get("codelist")
                if cl and cl.startswith("OMOP.") and cl in v1_omop:
                    f["codelist"] = cl
                    omop[cl] = v1_omop[cl]
            codes = OrderedDict()
            for line in desc.splitlines():
                if CODE_LINE.match(line):
                    for code, label in CODE_ITEM.findall(line):
                        codes[code] = label.strip()
            if codes:
                f["codelist"] = f"{t}.{name}"
                embedded[f["codelist"]] = codes
            for k, v in EXTRA.get((t, name), {}).items():
                f[k] = v
            fields.append(f)
        if pk is None:
            raise SystemExit(f"{t}: PK 가 없다")
        tables[t] = OrderedDict(name=t, kor=ws.title.split("(")[0], section=section, category=category(t), pk=pk,
                                fields=fields)
    embedded.update(EXTRA_CODELISTS)
    missing = [t for t in CORE + list(DISEASE) + EXTENSION + REFERENCE if t not in tables]
    if missing:
        raise SystemExit(f"수정안에 없는 테이블: {missing}")
    profiles = OrderedDict()
    for c, kor in DISEASE.items():
        tl = CORE + [c, "ADVERSE_EVENT"] + (["SACT"] if c in SACT_COHORTS else []) + REFERENCE
        profiles[c] = OrderedDict(kor=kor, tables=tl)

    os.makedirs(OUT, exist_ok=True)
    meta = OrderedDict(name="K-AI 질환별 코호트 명세 v3 (의뢰사 9/30 매뉴얼 + 2026-10-01 검토 결정)", version="3.0.0",
                       source=os.path.relpath(SRC, ROOT), generated_by="spec/build_spec_v3.py",
                       decisions="spec/client_v1/매뉴얼_수정_제안서_20261001.md", dropped_fields=dropped)

    class D(yaml.SafeDumper):
        pass
    D.add_representer(OrderedDict, lambda d, data: d.represent_dict(data.items()))
    D.ignore_aliases = lambda self, data: True
    dump = lambda obj, name: yaml.dump(obj, open(os.path.join(OUT, name), "w", encoding="utf-8"), Dumper=D,
                                       allow_unicode=True, sort_keys=False, width=160)
    dump(OrderedDict(meta=meta, tables=tables), "kai_cdm_spec_v3.yaml")
    dump(OrderedDict(embedded=embedded, omop=omop), "codelists_v3.yaml")
    dump(OrderedDict(cohorts=profiles), "cohort_profiles_v3.yaml")
    nf = sum(len(t["fields"]) for t in tables.values())
    print(f"v3 명세: 테이블 {len(tables)}개, 필드 {nf}개, 코드표 {len(embedded)}+{len(omop)}개, 제외 {dropped}")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""9월 30일 매뉴얼 수정 제안·의뢰사 확인 질문 검토 서버 (loopback 전용, 표준 라이브러리 + openpyxl)."""
import argparse
import csv
import io
import json
import os
import signal
import sqlite3
import sys
import threading
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

from proposals import CHANGES, LEGACY_LABELS, QUESTIONS

REVIEW_DIR = Path(__file__).resolve().parent
ROOT = REVIEW_DIR.parent
WORKBOOK = ROOT / "raw" / "매뉴얼_질환별 데이터 테이블_20260930.xlsx"

STATIC = {
    "/": ("index.html", "text/html; charset=utf-8"),
    "/index.html": ("index.html", "text/html; charset=utf-8"),
    "/app.js": ("app.js", "text/javascript; charset=utf-8"),
    "/style.css": ("style.css", "text/css; charset=utf-8"),
}

MAX_BODY = 64 * 1024
CHANGE_VERDICTS = ("adopt", "reject", "hold")
QUESTION_VERDICTS = ("answer", "hold")
VERDICT_LABEL = {"adopt": "채택", "reject": "채택 안 함", "hold": "보류", "answer": "답변"}
LIMITS = {"note": 4000, "reviewer": 100}


# ---------------------------------------------------------------- items

def _s(v):
    return "" if v is None else str(v)


def load_workbook_defs():
    """sheet title -> {field: {type, required, description, row}}"""
    import openpyxl
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheets = {}
    try:
        for ws in wb.worksheets:
            defs = {}
            for rno, row in enumerate(ws.iter_rows(values_only=True), 1):
                if rno <= 2 or not row or row[0] is None:
                    continue
                defs.setdefault(_s(row[0]).strip(), {
                    "type": _s(row[1]), "required": _s(row[3]), "description": _s(row[4]), "row": rno,
                })
            sheets[ws.title] = defs
    finally:
        wb.close()
    return sheets


def build_static():
    """제안·질문 정의를 워크북과 대조해 화면용 정적 항목으로 만든다."""
    sheets = load_workbook_defs()
    problems = []
    q_by_id = {q["id"]: q for q in QUESTIONS}
    changes = []
    for c in CHANGES:
        current = []
        for sheet, field in c["fields"]:
            d = sheets.get(sheet, {}).get(field)
            if d is None:
                problems.append(f"{c['id']}: 워크북 {sheet}에서 {field}를 찾지 못했습니다.")
                continue
            current.append({"sheet": sheet, "field": field, **d})
        for alt in c["requires"] or []:
            for qid, answers in alt.items():
                opts = {o["id"] for o in q_by_id[qid]["options"]} if qid in q_by_id else set()
                if not opts or not set(answers) <= opts:
                    problems.append(f"{c['id']}: 조건 {qid}={answers}가 질문 정의와 맞지 않습니다.")
        changes.append({
            "key": f"change:{c['id']}", "kind": "change", "id": c["id"], "rev": c.get("rev", ""), "title": c["title"],
            "priority": c["priority"], "sheet": c["sheet"], "current": current,
            "current_text": c.get("current", ""), "proposed": c["proposed"], "reason": c["reason"],
            "example": c["example"], "requires": c["requires"],
            "requires_text": requires_text(c["requires"], q_by_id),
        })
    questions = []
    for q in QUESTIONS:
        affects = [c["id"] for c in CHANGES if any(q["id"] in alt for alt in c["requires"] or [])]
        questions.append({
            "key": f"question:{q['id']}", "kind": "question", "id": q["id"], "rev": q.get("rev", ""), "title": q["title"],
            "why": q["why"], "options": q["options"], "recommended": q["recommended"],
            "legacy_keys": q["legacy"], "affects": affects,
        })
    return questions + changes, problems


def requires_text(requires, q_by_id):
    if not requires:
        return "항상 권장 (질문 답과 무관)"
    alts = []
    for alt in requires:
        parts = []
        for qid, answers in alt.items():
            labels = {o["id"]: o["label"] for o in q_by_id[qid]["options"]}
            parts.append(f"{qid} = " + " 또는 ".join(labels[a] for a in answers))
        alts.append(" 그리고 ".join(parts))
    return " / 또는 ".join(alts)


def need_status(requires, answers):
    """answers: {question id: answer id (답변한 질문만)} -> always | needed | not_needed | pending"""
    if not requires:
        return "always"
    pending = False
    for alt in requires:
        if all(answers.get(q) in a for q, a in alt.items()):
            return "needed"
        if not any(q in answers and answers[q] not in a for q, a in alt.items()):
            pending = True
    return "pending" if pending else "not_needed"


def conflict_of(status, decision):
    v = (decision or {}).get("verdict")
    if v == "adopt" and status == "not_needed":
        return "질문 답에 따르면 필요 없는 제안인데 채택했습니다."
    if v == "reject" and status == "needed":
        return "질문 답에 따르면 필요한 제안인데 채택하지 않았습니다. 해당 범위의 DQ는 수행할 수 없습니다."
    return None


# ---------------------------------------------------------------- storage

def default_db_path():
    if sys.platform == "darwin":
        base = Path.home() / "Library" / "Application Support"
    elif os.name == "nt":
        base = Path(os.environ.get("APPDATA") or Path.home() / "AppData" / "Roaming")
    else:
        base = Path(os.environ.get("XDG_DATA_HOME") or Path.home() / ".local" / "share")
    return base / "KAI_DQ_Review" / "review.db"


class Store:
    def __init__(self, path):
        self.path = Path(path).expanduser().resolve()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.Lock()
        self.db = sqlite3.connect(self.path, check_same_thread=False)
        self.db.row_factory = sqlite3.Row
        with self.lock, self.db:
            self.db.execute(
                "CREATE TABLE IF NOT EXISTS decisions ("
                "key TEXT PRIMARY KEY, verdict TEXT NOT NULL, note TEXT NOT NULL DEFAULT '',"
                "candidate_id TEXT NOT NULL DEFAULT '', candidate_label TEXT NOT NULL DEFAULT '',"
                "evidence TEXT NOT NULL DEFAULT '', reviewer TEXT NOT NULL DEFAULT '',"
                "updated_at TEXT NOT NULL)"
            )
            cols = {r["name"] for r in self.db.execute("PRAGMA table_info(decisions)")}
            if "resolution" not in cols:
                self.db.execute("ALTER TABLE decisions ADD COLUMN resolution TEXT NOT NULL DEFAULT ''")
            if "item_rev" not in cols:
                self.db.execute("ALTER TABLE decisions ADD COLUMN item_rev TEXT NOT NULL DEFAULT ''")

    def all(self):
        with self.lock:
            rows = self.db.execute("SELECT key, verdict, resolution, note, reviewer, updated_at, item_rev FROM decisions").fetchall()
        return {r["key"]: {k: r[k] for k in ("verdict", "resolution", "note", "reviewer", "updated_at", "item_rev")} for r in rows}

    def upsert(self, key, d):
        ts = datetime.now(timezone.utc).isoformat(timespec="seconds")
        with self.lock, self.db:
            self.db.execute(
                "INSERT INTO decisions(key,verdict,resolution,note,reviewer,updated_at,item_rev) VALUES(?,?,?,?,?,?,?)"
                " ON CONFLICT(key) DO UPDATE SET verdict=excluded.verdict, resolution=excluded.resolution,"
                " note=excluded.note, reviewer=excluded.reviewer, updated_at=excluded.updated_at, item_rev=excluded.item_rev",
                (key, d["verdict"], d["resolution"], d["note"], d["reviewer"], ts, d["item_rev"]))
        return {**d, "updated_at": ts}

    def close(self):
        with self.lock:
            self.db.close()


class ApiError(Exception):
    def __init__(self, status, message):
        super().__init__(message)
        self.status, self.message = status, message


def validate_decision(payload, items_by_key):
    if not isinstance(payload, dict):
        raise ApiError(400, "JSON 객체가 필요합니다.")
    extra = set(payload) - {"key", "verdict", "resolution", "note", "reviewer"}
    if extra:
        raise ApiError(400, f"허용되지 않는 필드: {', '.join(sorted(extra))}")
    item = items_by_key.get(payload.get("key"))
    if item is None:
        raise ApiError(400, "알 수 없는 key입니다.")
    verdict, resolution = payload.get("verdict"), payload.get("resolution", "")
    if not isinstance(resolution, str):
        raise ApiError(400, "resolution은 문자열이어야 합니다.")
    if item["kind"] == "change":
        if verdict not in CHANGE_VERDICTS:
            raise ApiError(400, "제안 결정은 adopt, reject, hold 중 하나여야 합니다.")
        if resolution:
            raise ApiError(400, "제안에는 답변 선택지가 없습니다.")
    else:
        if verdict not in QUESTION_VERDICTS:
            raise ApiError(400, "질문 결정은 answer 또는 hold여야 합니다.")
        options = {o["id"] for o in item["options"]}
        if verdict == "answer" and resolution not in options:
            raise ApiError(400, "질문의 선택지 중 하나를 고르세요.")
        if verdict == "hold" and resolution:
            raise ApiError(400, "보류에는 답을 함께 저장하지 않습니다.")
    out = {"verdict": verdict, "resolution": resolution}
    for f, limit in LIMITS.items():
        v = payload.get(f, "") or ""
        if not isinstance(v, str):
            raise ApiError(400, f"{f}는 문자열이어야 합니다.")
        v = v.strip()
        if len(v) > limit:
            raise ApiError(400, f"{f}는 {limit}자 이하여야 합니다.")
        out[f] = v
    out["item_rev"] = item["rev"]
    return item["key"], out


# ---------------------------------------------------------------- exports

def _csv_safe(v):
    v = _s(v)
    return "'" + v if v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


def _current_text(it):
    if it["current_text"]:
        return it["current_text"]
    return "\n".join(f"{c['sheet']} {c['row']}행 {c['field']} | {c['type']} | 필수 {c['required'] or '(빈칸)'} | {c['description']}"
                     for c in it["current"])


def _answer_label(it):
    d = it["decision"] or {}
    return next((o["label"] for o in it["options"] if o["id"] == d.get("resolution")), "")


def to_csv(items):
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow(["key", "kind", "id", "title", "priority", "sheet", "current", "proposed", "reason", "example",
                "condition", "need", "conflict", "decision", "answer", "note", "reviewer", "updated_at"])
    for it in items:
        d = it["decision"] or {}
        if it["kind"] == "change":
            row = [it["priority"], it["sheet"], _current_text(it), it["proposed"], it["reason"], it["example"],
                   it["requires_text"], it["need"], it["conflict"],
                   VERDICT_LABEL.get(d.get("verdict"), "재결정 필요" if it["previous_decision"] else ""), ""]
        else:
            row = ["", "", "", "", it["why"], "", "", "", "",
                   VERDICT_LABEL.get(d.get("verdict"), "재결정 필요" if it["previous_decision"] else ""), _answer_label(it)]
        w.writerow([it["key"], it["kind"], it["id"], it["title"]] + [_csv_safe(v) for v in row]
                   + [_csv_safe(d.get("note")), _csv_safe(d.get("reviewer")), _s(d.get("updated_at"))])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


NEED_LABEL = {"always": "항상 권장", "needed": "필요 (질문 답 기준)", "not_needed": "불필요 (질문 답 기준)",
              "pending": "질문 답 대기"}


def to_markdown(items, source_name):
    lines = [f"# 매뉴얼 수정 제안서 — {source_name}", "",
             f"작성: {datetime.now(timezone.utc).date().isoformat()} · 검토 웹서비스 내보내기", ""]
    lines += ["## 1. 의뢰사 확인 질문", ""]
    for it in (i for i in items if i["kind"] == "question"):
        d = it["decision"] or {}
        state = f"답변: **{_answer_label(it)}**" if d.get("verdict") == "answer" else ("보류" if d else "미답변")
        lines += [f"### {it['id']}. {it['title']}", "", it["why"], "",
                  *[f"- {o['label']} → {o['effect']}" for o in it["options"]], "", f"상태: {state}"]
        if d.get("note"):
            lines.append(f"메모: {d['note']}")
        lines.append("")
    lines += ["## 2. 매뉴얼 수정 제안", ""]
    for it in (i for i in items if i["kind"] == "change"):
        d = it["decision"] or {}
        lines += [f"### {it['id']}. {it['title']} ({it['priority']})", "",
                  f"- 적용 조건: {it['requires_text']} → **{NEED_LABEL[it['need']]}**",
                  f"- 검토 결정: {VERDICT_LABEL.get(d.get('verdict'), '재결정 필요 (문구 변경)' if it['previous_decision'] else '미결정')}", "",
                  "현재:", "```", _current_text(it), "```", "제안:", "```", it["proposed"], "```",
                  f"이유: {it['reason']}", ""]
        if it["example"]:
            lines += [f"예시: {it['example']}", ""]
        if it["conflict"]:
            lines += [f"⚠ 모순: {it['conflict']}", ""]
        if d.get("note"):
            lines += [f"메모: {d['note']}", ""]
    return "\n".join(lines).encode("utf-8")


# ---------------------------------------------------------------- http

class App:
    def __init__(self, store, port):
        self.store = store
        self.port = port
        self.items, self.problems = build_static()
        self.items_by_key = {i["key"]: i for i in self.items}
        self.source_name = WORKBOOK.name
        self.warning = "" if not self.problems else "정의 불일치: " + "; ".join(self.problems)

    def payload(self):
        dec = self.store.all()
        # 문구가 바뀐 항목(rev 불일치)의 이전 결정은 유효한 결정으로 쓰지 않고 참고로만 돌려준다.
        current, previous = {}, {}
        for it in self.items:
            d = dec.get(it["key"])
            if d is None:
                continue
            (current if d["item_rev"] == it["rev"] else previous)[it["key"]] = d
        answers = {}
        for it in self.items:
            d = current.get(it["key"])
            if it["kind"] == "question" and d and d["verdict"] == "answer":
                answers[it["id"]] = d["resolution"]
        out = []
        for it in self.items:
            d = current.get(it["key"])
            row = {**it, "decision": d, "previous_decision": previous.get(it["key"])}
            if it["kind"] == "change":
                row["need"] = need_status(it["requires"], answers)
                row["conflict"] = conflict_of(row["need"], d)
            else:
                row["legacy"] = [
                    {"key": k, "verdict": dec[k]["verdict"], "label": LEGACY_LABELS.get(dec[k]["resolution"], dec[k]["resolution"]),
                     "note": dec[k]["note"], "reviewer": dec[k]["reviewer"]}
                    for k in it["legacy_keys"] if k in dec
                ]
            out.append(row)
        return {"items": out, "source_name": self.source_name, "warning": self.warning}

    def allowed_hosts(self):
        return {f"127.0.0.1:{self.port}", f"localhost:{self.port}", f"[::1]:{self.port}"}


def make_handler(app):
    class Handler(BaseHTTPRequestHandler):
        server_version = "KAIReview/2"
        protocol_version = "HTTP/1.1"

        def log_message(self, fmt, *args):
            sys.stderr.write("%s - %s\n" % (self.address_string(), fmt % args))

        def _send(self, status, body, ctype, extra=None):
            self.send_response(status)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Referrer-Policy", "no-referrer")
            self.send_header("Content-Security-Policy", "default-src 'self'; frame-ancestors 'none'")
            for k, v in (extra or {}).items():
                self.send_header(k, v)
            self.end_headers()
            if self.command != "HEAD":
                self.wfile.write(body)

        def _json(self, status, obj, extra=None):
            self._send(status, json.dumps(obj, ensure_ascii=False).encode("utf-8"),
                       "application/json; charset=utf-8", extra)

        def _error(self, status, message):
            self._json(status, {"error": message}, {"Connection": "close"} if status >= 400 else None)
            self.close_connection = True

        def _check_host(self):
            if (self.headers.get("Host") or "").lower() not in app.allowed_hosts():
                raise ApiError(403, "허용되지 않은 Host입니다.")

        def _check_origin(self):
            origin = self.headers.get("Origin")
            if origin is None:
                return
            if urlsplit(origin).scheme != "http" or urlsplit(origin).netloc.lower() not in app.allowed_hosts():
                raise ApiError(403, "허용되지 않은 Origin입니다.")

        def _dispatch(self, fn):
            try:
                self._check_host()
                fn()
            except ApiError as e:
                self._error(e.status, e.message)
            except (BrokenPipeError, ConnectionResetError):
                self.close_connection = True
            except Exception as e:  # noqa: BLE001
                sys.stderr.write(f"internal error: {e!r}\n")
                self._error(500, "서버 내부 오류입니다.")

        def do_GET(self):
            self._dispatch(self._get)

        do_HEAD = do_GET

        def do_POST(self):
            self._dispatch(self._post)

        def _get(self):
            path = urlsplit(self.path).path
            if path in STATIC:
                name, ctype = STATIC[path]
                f = REVIEW_DIR / name
                if not f.is_file():
                    raise ApiError(404, f"{name} 파일이 없습니다.")
                self._send(200, f.read_bytes(), ctype)
            elif path == "/api/review":
                self._json(200, app.payload())
            elif path == "/api/export.json":
                data = app.payload()
                data["exported_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
                self._json(200, data, {"Content-Disposition": 'attachment; filename="review_export.json"'})
            elif path == "/api/export.csv":
                self._send(200, to_csv(app.payload()["items"]), "text/csv; charset=utf-8",
                           {"Content-Disposition": 'attachment; filename="review_export.csv"'})
            elif path == "/api/export.md":
                self._send(200, to_markdown(app.payload()["items"], app.source_name), "text/markdown; charset=utf-8",
                           {"Content-Disposition": 'attachment; filename="manual_change_proposal.md"'})
            else:
                raise ApiError(404, "찾을 수 없습니다.")

        def _post(self):
            if urlsplit(self.path).path != "/api/decisions":
                raise ApiError(404, "찾을 수 없습니다.")
            self._check_origin()
            if (self.headers.get("Content-Type") or "").split(";")[0].strip().lower() != "application/json":
                raise ApiError(415, "Content-Type은 application/json이어야 합니다.")
            if "chunked" in (self.headers.get("Transfer-Encoding") or "").lower():
                raise ApiError(411, "Content-Length가 필요합니다.")
            try:
                n = int(self.headers.get("Content-Length", ""))
            except ValueError:
                raise ApiError(411, "Content-Length가 필요합니다.") from None
            if n < 0:
                raise ApiError(400, "잘못된 Content-Length입니다.")
            if n > MAX_BODY:
                raise ApiError(413, "요청 본문이 너무 큽니다.")
            try:
                payload = json.loads(self.rfile.read(n).decode("utf-8"))
            except (UnicodeDecodeError, ValueError):
                raise ApiError(400, "올바른 JSON이 아닙니다.") from None
            key, dec = validate_decision(payload, app.items_by_key)
            app.store.upsert(key, dec)
            self._json(200, app.payload())

    return Handler


def main():
    ap = argparse.ArgumentParser(description="매뉴얼 수정 제안·의뢰사 질문 검토 서버")
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--db", default=None, help=f"SQLite 경로 (기본: {default_db_path()})")
    args = ap.parse_args()
    if not 0 <= args.port <= 65535:
        ap.error("--port 범위 오류")

    store = Store(args.db or default_db_path())
    httpd = ThreadingHTTPServer(("127.0.0.1", args.port), lambda *a, **k: None)
    port = httpd.server_address[1]
    try:
        app = App(store, port)
    except Exception:
        httpd.server_close()
        store.close()
        raise
    httpd.RequestHandlerClass = make_handler(app)
    httpd.daemon_threads = True

    def stop(*_):
        threading.Thread(target=httpd.shutdown, daemon=True).start()

    signal.signal(signal.SIGINT, stop)
    signal.signal(signal.SIGTERM, stop)
    print(f"검토 서버: http://127.0.0.1:{port}/  (DB: {store.path})", flush=True)
    if app.problems:
        print("경고:", *app.problems, sep="\n  ", file=sys.stderr)
    try:
        httpd.serve_forever()
    finally:
        httpd.server_close()
        store.close()
        print("종료", flush=True)


if __name__ == "__main__":
    main()

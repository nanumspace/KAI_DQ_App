# -*- coding: utf-8 -*-
"""
사전 v2 변환기 · 배치표(엑셀) → 저장 모델(OMOP 5.4 + Oncology 구조 + K-AI 확장), CRF 서식, K-AI 어휘, 서식→레코드 매핑과 예시.

입력
  internal/decisions/배치표_v0_YYYYMMDD.xlsx    시트: 자문항목, 현명세필드, 테이블정의, 코드표
  spec/kai_cdm_spec.yaml, spec/codelists.yaml   v1 (유지 테이블 ANTP_THERAPY·ADVERSE_EVENT 의 필드와 내장 코드표)
  spec/v2_field_names.py                         자문 항목 → 영문 필드명·단위·코드표·타입
  spec/v2_omop_map.py                            OMOP 테이블 정의, 어휘 씨앗, 서식→레코드 규칙
  spec/v2/vocab/concept_registry.json            (있으면) 기존 concept_id 보존

출력 (spec/v2/)
  kai_cdm_spec_v2.yaml        저장 테이블: OMOP 14개 + K-AI 확장(병리·바이오마커·질환)
  crf_forms_v2.yaml           병원이 채우는 서식과 필드별 저장 위치
  crf_to_omop_examples.md     서식 한 행이 어떤 레코드가 되는지 (폐암 환자 예시)
  vocab/CONCEPT.csv, VOCABULARY.csv, CONCEPT_RELATIONSHIP.csv, CONCEPT_SET.csv, CONCEPT_SET_ITEM.csv, concept_registry.json
  derived_v2.yaml             코어에서 파생하는 변수와 파생 점수
  legacy_fields_v2.yaml       v1 필드 중 합쳐지지 않은 것
  dictionary_v2.xlsx          팀 배포용 사전
  spec_v2_report.md           통계·검토 목록

    python spec/build_spec_v2.py [배치표.xlsx] [--out spec/v2]
"""
import sys, os, re, json, datetime, collections, csv
import yaml, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "spec"))
from v2_field_names import FIELD_NAMES, UNITS, CODELIST_OVERRIDE, TYPE_OVERRIDE  # noqa: E402
from v2_omop_map import CONCEPT_ID_BASE, OMOP_TABLES, SEED_CONCEPTS, FORM_MAP  # noqa: E402

args = [a for a in sys.argv[1:] if not a.startswith("--")]
XLSX = args[0] if args else os.path.join(ROOT, "internal/decisions/배치표_v0_20260912.xlsx")
OUT = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else os.path.join(ROOT, "spec/v2")
VOCAB = os.path.join(OUT, "vocab"); os.makedirs(VOCAB, exist_ok=True)
TODAY = datetime.date.today().isoformat()

V1 = yaml.safe_load(open(os.path.join(ROOT, "spec/kai_cdm_spec.yaml"), encoding="utf-8"))["tables"]
V1_CODES = yaml.safe_load(open(os.path.join(ROOT, "spec/codelists.yaml"), encoding="utf-8"))
COHORT_OF = {"폐암": "LUNG_CANCER", "유방암": "BREAST_CANCER", "대장암": "COLORECTAL_CANCER", "대사이상지방간염": "MASLD", "당뇨병": "DIABETES", "림프종": "LYMPHOMA"}
COHORT_KOR = {"LUNG_CANCER": "폐암", "BREAST_CANCER": "유방암", "COLORECTAL_CANCER": "대장암", "MASLD": "대사이상지방간", "DIABETES": "당뇨병", "LYMPHOMA": "림프종"}
ALL_COHORTS = list(COHORT_OF.values())
DISEASE_CRF = {"LUNG_CRF": "LUNG", "BREAST_CRF": "BREAST", "COLORECTAL_CRF": "COLORECTAL", "MASLD_CRF": "MASLD", "DIABETES_CRF": "DIABETES", "LYMPHOMA_CRF": "LYMPHOMA"}
BASE_COHORT = {"LUNG": "LUNG_CANCER", "BREAST": "BREAST_CANCER", "COLORECTAL": "COLORECTAL_CANCER", "MASLD": "MASLD", "DIABETES": "DIABETES", "LYMPHOMA": "LYMPHOMA"}
KEEP_V1 = {"ANTP_THERAPY", "ADVERSE_EVENT"}
V1_ALIAS = {"line_number": "regimen_or_line_of_therapy", "regimen_name": "regimen", "cycles_count": "number_of_cycles"}   # 자문 항목 → v1 유지 필드
AE_DROP = {"observation_id", "condition_occurrence_id", "measurement_id", "procedure_occurrence_id", "drug_exposure_id", "note_id"}  # 연결은 EPISODE_EVENT 로
warnings = []

# ================================================================ 1. 배치표 → 서식(form)
wb = openpyxl.load_workbook(XLSX, data_only=True)
def rows_of(sheet):
    ws = wb[sheet]; hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
ITEMS, V1FIELDS, TDEF, CODES = rows_of("자문항목"), rows_of("현명세필드"), rows_of("테이블정의"), rows_of("코드표")
CODE_IDS = {c["코드표"] for c in CODES}

forms = collections.OrderedDict()
def add_form(name, kor, grain, pk, parent, desc):
    forms[name] = dict(name=name, kor=kor, grain=grain, pk=pk, parent=parent or None, desc=desc, cohorts=set(), fields=collections.OrderedDict())
for t in TDEF:
    if t["층"] == "코어": continue
    add_form(t["테이블"], t["한글"], t["행 단위"], t["고유키"], t["상위 FK"], t["설명"])
SPLIT = {"환자 1": ("BASELINE", "환자 1", "기저(환자 1)", "COHORT"), "검사 1": ("EXAM", "검사 1", "검사(검사 1)", "PROCEDURE_OCCURRENCE"), "사건 1": ("EVENT", "사건 1", "사건(사건 1)", "COHORT"),
         "평가 1": ("ASSESSMENT", "평가 1", "평가(평가 1)", "COHORT"), "수술 1": ("SURGERY_DETAIL", "수술 1 (SURGERY 서식 자식)", "수술 세부(수술 1)", "PROCEDURE_OCCURRENCE")}
def split_disease(dest, grain_text):
    base = DISEASE_CRF[dest]
    key = next((k for k in SPLIT if grain_text and grain_text.startswith(k)), "환자 1")
    suffix, grain, kor, parent = SPLIT[key]; name = f"{base}_{suffix}"
    if name not in forms: add_form(name, f"{COHORT_KOR[BASE_COHORT[base]]} {kor}", grain, f"{name.lower()}_id", parent, forms[dest]["desc"] + f" — 행 단위 {grain}")
    return name

def field_type(kind, name):
    if kind == "날짜": return "date"
    if kind in ("예/아니오", "예/아니오+날짜"): return "integer"
    if kind == "수치": return "integer" if re.search(r"count|_n$|number|regions|sites_count|children|fractions|cycles|lines|line_number|age$|year$|score$|grade$", name) and not re.search(r"percent|ratio", name) else "float"
    if kind == "텍스트": return "text"
    return "varchar"
def unit_of(name, item):
    m = re.search(r"\((Gy|cm|mm|mg/dL|%|kPa|dB/m|μm|um|mmHg|mL|ng/mL)\)", item)
    return ({"μm": "um", "mmHg": "mm[Hg]"}.get(m.group(1), m.group(1)) if m else UNITS.get(name))
def codelist_for(row, name, kind):
    if name in CODELIST_OVERRIDE: return CODELIST_OVERRIDE[name]
    cl = (row.get("코드표") or "").strip()
    if kind in ("예/아니오", "예/아니오+날짜"): return "YN"
    if kind != "닫힌 범주": return None
    if not cl or cl.endswith("*") or "/" in cl:
        first = cl.split("/")[0].strip() if cl else ""
        return first if first and not first.endswith("*") and first in CODE_IDS else "CL_" + name.upper()
    return cl
def tier_of(g): g = str(g or ""); return "필수" if g.startswith("필수") else ("권고" if g.startswith("권고") else "보류")
def new_field(name, kor, ty, kind, tier, **kw):
    f = dict(name=name, kor=kor, type=ty, value_kind=kind, tier=tier, required=False, codelist=None, unit=None, standard=None, pattern=None, vocabulary=None, cohorts=set(), also=[], v1_hint=[], desc=None)
    f.update(kw); return f

derived, scores = [], []
for r in ITEMS:
    cohort, item, way, dest, kind = COHORT_OF[r["질환"]], r["항목"], r["갈래"], r["목적지 테이블"], (r["값 종류"] or "닫힌 범주")
    name = FIELD_NAMES.get(item) or ("item_" + re.sub(r"\W+", "_", item)[:30].lower())
    if item not in FIELD_NAMES: warnings.append(f"영문 필드명 없음: {item}")
    tier, std = tier_of(r["등급"]), (r.get("표준 매핑 후보") or "")
    if way == "코어 파생":
        derived.append(dict(name=name, kor=item, cohort=cohort, source_table=dest, value_kind=kind, tier=tier, concept_set=(r.get("코드표") or None), standard=std or None, note=r.get("메모") or None)); continue
    if dest in DISEASE_CRF: dest = split_disease(dest, r.get("행 단위"))
    if dest not in forms: warnings.append(f"목적지 미정의: {dest} ← {item}"); continue
    if dest in KEEP_V1 and name in V1_ALIAS: name = V1_ALIAS[name]
    F = forms[dest]; F["cohorts"].add(cohort)
    f = F["fields"].get(name)
    if f is None:
        ty, pat = TYPE_OVERRIDE.get(name, (field_type(kind, name), None))
        cl = codelist_for(r, name, kind)
        if cl == "YN": kind = "예/아니오"
        if pat: kind = "패턴"
        elif ty in ("float", "integer") and kind == "닫힌 범주": kind = "수치"
        f = new_field(name, item, ty, kind, tier, codelist=cl, unit=unit_of(name, item) if ty in ("float", "integer") else None, standard=std or None, pattern=pat,
                      vocabulary=(std or "참조표") if kind == "열린 표준 코드" else None, derived_score=(way == "파생 점수"))
        if way == "파생 점수": scores.append(dict(name=name, kor=item, form=dest, cohort=cohort))
        F["fields"][name] = f
    else:
        if f["kor"] != item and item not in f["also"]: f["also"].append(item)
        if tier == "필수": f["tier"] = "필수"
        elif tier == "권고" and f["tier"] == "보류": f["tier"] = "권고"
    f["cohorts"].add(cohort)
    if r.get("현 명세 대응 필드") and f"{r['현 테이블']}.{r['현 명세 대응 필드']}" not in f["v1_hint"]: f["v1_hint"].append(f"{r['현 테이블']}.{r['현 명세 대응 필드']}")
    if r.get("메모") and r["메모"] not in (f["desc"] or ""): f["desc"] = ((f["desc"] + "; ") if f["desc"] else "") + r["메모"]
for d in list(DISEASE_CRF):
    if d in forms and not forms[d]["fields"]: del forms[d]

# v1 유지 서식(ANTP·AE)의 나머지 필드
for t in KEEP_V1:
    F = forms[t]
    for vf in V1[t]["fields"]:
        n = vf["name"]
        if n in F["fields"] or (t == "ADVERSE_EVENT" and n in AE_DROP): continue
        kind = "닫힌 범주" if vf.get("codelist") else ("날짜" if vf.get("type") == "date" else ("수치" if vf.get("type") in ("integer", "float") else ("열린 표준 코드" if n.endswith("_concept_id") else "텍스트")))
        cl_v1 = vf.get("codelist"); cl_v1 = "YN" if cl_v1 and str(cl_v1).startswith("common") else cl_v1
        if cl_v1 == "YN": kind = "예/아니오"
        F["fields"][n] = new_field(n, vf.get("kor") or n, vf.get("type", "varchar"), kind, "필수" if vf.get("required") else "권고", codelist=cl_v1, v1_hint=[f"{t}.{n}"], desc=(vf.get("desc") or "").replace("\n", " ")[:160])
    F["cohorts"] |= set(ALL_COHORTS) if t == "ADVERSE_EVENT" else {"LUNG_CANCER", "BREAST_CANCER", "COLORECTAL_CANCER", "LYMPHOMA"}
    for f in F["fields"].values():
        if not f["cohorts"]: f["cohorts"] = set(F["cohorts"])
ae = forms["ADVERSE_EVENT"]["fields"]
for n, k, cl, kind in [("adverse_event_concept_id", "이상사례 용어(MedDRA PT)", None, "열린 표준 코드"), ("severity_concept_id", "CTCAE 등급", "CTCAE_GRADE", "닫힌 범주"), ("causality_concept_id", "인과성", "AE_CAUSALITY", "닫힌 범주"), ("seriousness_concept_id", "중대성", "AE_SERIOUSNESS", "닫힌 범주")]:
    if n in ae: ae[n].update(kor=k, codelist=cl or ae[n]["codelist"], value_kind=kind, vocabulary="MedDRA" if n == "adverse_event_concept_id" else None)
ae["adverse_event_source_value"] = new_field("adverse_event_source_value", "이상사례 원천 표기", "varchar", "텍스트", "권고", cohorts=set(forms["ADVERSE_EVENT"]["cohorts"]), desc="병원 표기(MedDRA 미보유 시 CTCAE 용어)")
fh = forms["FAMILY_HISTORY"]["fields"]
fh["family_disease"] = new_field("family_disease", "가족 질환(암종 외)", "varchar", "열린 표준 코드", "권고", vocabulary="SNOMED/KCD", cohorts=set(forms["FAMILY_HISTORY"]["cohorts"]))
fh["family_member_age_at_diagnosis"] = new_field("family_member_age_at_diagnosis", "가족 진단 나이", "integer", "수치", "권고", unit="a", cohorts=set(forms["FAMILY_HISTORY"]["cohorts"]))
if "antp_end_date" not in forms["ANTP_THERAPY"]["fields"]: pass

# 구조 필드 (서식 키·사건일)
EVENT_DATE = {"COHORT_INDEX": ("index_date", "코호트 진입일(기준일)"), "STAGING": ("staging_date", "병기 평가일"), "PATHOLOGY_REPORT": ("specimen_date", "검체 채취일"), "BIOMARKER": ("test_date", "검사일"),
              "SURGERY": ("surgery_date", "수술 시행일"), "RADIOTHERAPY": ("rt_start_date", "방사선치료 시작일"), "RESPONSE_ASSESSMENT": ("response_assessment_date", "반응평가 기준일"),
              "FAMILY_HISTORY": ("record_date", "기록일"), "PATHOLOGY_TUMOR": (None, None), "ANTP_THERAPY": ("antp_start_date", "항암요법 시작일"), "ADVERSE_EVENT": ("adverse_event_start_date", "발생일")}
def structural(F):
    name, pk, fields = F["name"], F["pk"], collections.OrderedDict()
    fields[pk] = new_field(pk, f"{F['kor']} ID", "bigint", "키", "필수", required=True, key="pk", cohorts=set(F["cohorts"]), desc="행 고유키")
    if pk != "person_id": fields["person_id"] = new_field("person_id", "환자 ID", "bigint", "키", "필수", required=True, key="fk", ref="PERSON.person_id", cohorts=set(F["cohorts"]))
    if name == "PATHOLOGY_REPORT": fields["specimen_id"] = new_field("specimen_id", "검체 ID", "bigint", "키", "필수", required=True, key="fk", ref="SPECIMEN.specimen_id", cohorts=set(F["cohorts"]), desc="SPECIMEN 레코드(검체 종류·부위·채취일)")
    if name == "PATHOLOGY_TUMOR": fields["pathology_report_id"] = new_field("pathology_report_id", "병리 보고서 ID", "bigint", "키", "필수", required=True, key="fk", ref="PATHOLOGY_REPORT.pathology_report_id", cohorts=set(F["cohorts"]))
    if name == "BIOMARKER": fields["specimen_id"] = new_field("specimen_id", "검체 ID", "bigint", "키", "권고", key="fk", ref="SPECIMEN.specimen_id", cohorts=set(F["cohorts"]), desc="조직 기반 검사이면 검체 연결")
    if name.endswith("_SURGERY_DETAIL"): fields["procedure_occurrence_id"] = new_field("procedure_occurrence_id", "수술 시술 ID", "bigint", "키", "필수", required=True, key="fk", ref="PROCEDURE_OCCURRENCE.procedure_occurrence_id", cohorts=set(F["cohorts"]))
    if name in ("RESPONSE_ASSESSMENT", "ADVERSE_EVENT") and "antp_id" not in F["fields"]: fields["antp_id"] = new_field("antp_id", "관련 항암요법 ID", "bigint", "키", "권고", key="fk", ref="ANTP_THERAPY.antp_id", cohorts=set(F["cohorts"]))
    if name not in ("FAMILY_HISTORY", "PATHOLOGY_TUMOR", "COHORT_INDEX"): fields["visit_occurrence_id"] = new_field("visit_occurrence_id", "방문 ID", "bigint", "키", "권고", key="fk", ref="VISIT_OCCURRENCE.visit_occurrence_id", cohorts=set(F["cohorts"]))
    ed = EVENT_DATE.get(name) or {"BASELINE": ("record_date", "기록일"), "EXAM": ("exam_date", "검사일"), "EVENT": ("event_date", "사건 발생일"), "ASSESSMENT": ("assessment_date", "평가일"), "SURGERY_DETAIL": (None, None)}[name.split("_", 1)[1]]
    if ed and ed[0]:
        if ed[0] in F["fields"]: F["fields"][ed[0]].update(required=True, event_date=True)
        else: fields[ed[0]] = new_field(ed[0], ed[1], "date", "날짜", "필수", required=True, event_date=True, cohorts=set(F["cohorts"]), desc="행의 기준 사건일")
    for k, v in F["fields"].items():
        if k not in fields: fields[k] = v
    F["fields"] = fields
for F in forms.values(): structural(F)
OMOP_FORMS = [f for f in FORM_MAP if f in forms]
EXT_FORMS = [f for f in forms if f not in FORM_MAP]

# ================================================================ 2. 코드표 → 값 집합
codelists = collections.OrderedDict()
for c in CODES:
    cid = c["코드표"]; cl = codelists.setdefault(cid, dict(kor=c["한글"], standard=None, note=None, values=[]))
    if c.get("표준 매핑 후보(코드표 단위)"): cl["standard"] = c["표준 매핑 후보(코드표 단위)"]
    if c.get("비고"): cl["note"] = c["비고"]
    val = str(c["값(코드·라벨)"]).strip(); code, _, label = val.partition(" ")
    if not label or not re.match(r"^([0-9]+|[A-Z0-9+\-/.]{1,5}|[a-z]{1,3})$", code): code, label = re.sub(r"\s+", "_", val), val
    cl["values"].append(dict(code=code, label=label, snomed=(str(c["SNOMED/표준 ID"]).strip() if c.get("SNOMED/표준 ID") else None)))
for k, v in V1_CODES.get("embedded", {}).items():
    t, f = k.split(".", 1)
    if t in KEEP_V1: codelists[f"{t}.{f}"] = dict(kor=f, standard=None, note="v1 내장 코드표", values=[dict(code=str(a), label=str(b), snomed=None) for a, b in v.items()])
for F in forms.values():
    for f in F["fields"].values():
        if f.get("codelist") and "." in str(f["codelist"]) and f["codelist"] not in codelists: f["codelist"] = "CL_" + f["name"].upper()
placeholder = sorted({f["codelist"] for F in forms.values() for f in F["fields"].values() if f.get("codelist") and str(f["codelist"]).startswith("CL_")})
for p in placeholder: codelists[p] = dict(kor=p, standard=None, note="값 집합 미정의 · 2.1 에서 채움", values=[])

# ================================================================ 3. K-AI 어휘 (concept_id ≥ 10,000,000,000, 등록부로 고정)
reg_path = os.path.join(VOCAB, "concept_registry.json")
registry = json.load(open(reg_path, encoding="utf-8")) if os.path.exists(reg_path) else {}
next_id = max([CONCEPT_ID_BASE] + list(registry.values())) + 1
concepts = collections.OrderedDict()
def concept(key, ko, en, domain, vocab, cls, code, standard="S", hint=None):
    global next_id
    if key not in registry: registry[key] = next_id; next_id += 1
    cid = registry[key]
    if cid not in concepts: concepts[cid] = dict(concept_id=cid, key=key, concept_name_ko=ko, concept_name_en=en or ko, domain_id=domain, vocabulary_id=vocab, concept_class_id=cls, concept_code=code, standard_concept=standard, standard_hint=hint or "", valid_start_date="2026-09-12", valid_end_date="2099-12-31", invalid_reason="")
    return cid
for key, (ko, en, dom, voc, cls, code, std) in SEED_CONCEPTS.items(): concept(key, ko, en, dom, voc, cls, code, std)
for ucum in sorted({f["unit"] for F in forms.values() for f in F["fields"].values() if f.get("unit")}): concept(f"UNIT:{ucum}", ucum, ucum, "Unit", "UCUM", "Unit", ucum)
CS_DOMAIN = {"STAGE": "Modifier", "TNM": "Modifier", "ECOG": "Observation", "RESPONSE": "Measurement", "DEAUVILLE": "Measurement", "RCB": "Modifier", "CTCAE": "Measurement", "AE_": "Observation", "YN": "Meas Value",
             "FAMILY": "Observation", "DISCHARGE": "Visit", "SMOKING": "Observation", "SURG": "Procedure", "RT_": "Procedure", "BIOMARKER": "Meas Value", "BIRADS": "Measurement", "FIBROSIS": "Modifier", "NAS": "Modifier",
             "CHILD": "Measurement", "CKD": "Condition", "DR_": "Condition", "WAGNER": "Condition", "DM_": "Condition", "RISK": "Measurement", "SCT": "Procedure"}
concept_sets = collections.OrderedDict()
for cid_, cl in codelists.items():
    dom = next((d for p, d in CS_DOMAIN.items() if cid_.startswith(p)), "Meas Value")
    parent = concept(f"CS:{cid_}", cl["kor"], cid_, dom, "KAI", "Value set", f"KAI-CS-{cid_}", "C", cl.get("standard"))
    concept_sets[cid_] = dict(concept_set_id=cid_, kor=cl["kor"], parent_concept_id=parent, note=cl.get("note"), items=[])
    for v in cl["values"]:
        voc, code = ("SNOMED", v["snomed"]) if v.get("snomed") else ("KAI", f"KAI-{cid_}-{v['code']}")
        c = concept(f"CS:{cid_}:{v['code']}", v["label"], v["label"], dom, voc, "Value", code, "S")
        concept_sets[cid_]["items"].append(dict(concept_id=c, code=v["code"], label=v["label"]))
for F in forms.values():
    for f in F["fields"].values():
        if f.get("key") or f.get("event_date"): continue
        dom = "Measurement" if f["value_kind"] == "수치" or (F["name"] in FORM_MAP and FORM_MAP[F["name"]].get("attr_default") == "MEASUREMENT") else "Observation"
        f["concept_id"] = concept(f"FIELD:{F['name']}:{f['name']}", f["kor"], f["name"], dom, "KAI", "Attribute", f"KAI-{F['name']}-{f['name']}", "S", f.get("standard"))
for d in derived:
    cs = d.get("concept_set")
    if cs and cs not in concept_sets:
        parent = concept(f"CS:{cs}", cs, cs, "Condition" if "COMORBID" in cs else ("Drug" if "DRUG" in cs else "Observation"), "KAI", "Value set", f"KAI-CS-{cs}", "C")
        concept_sets[cs] = dict(concept_set_id=cs, kor=cs, parent_concept_id=parent, note="파생 변수용 concept 목록 · 2.1 에서 채움", items=[])
json.dump(registry, open(reg_path, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
def wcsv(fn, rows, cols):
    with open(os.path.join(VOCAB, fn), "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols); w.writeheader()
        for r in rows: w.writerow({c: r.get(c, "") for c in cols})
wcsv("CONCEPT.csv", concepts.values(), ["concept_id", "concept_name_ko", "concept_name_en", "domain_id", "vocabulary_id", "concept_class_id", "concept_code", "standard_concept", "standard_hint", "valid_start_date", "valid_end_date", "invalid_reason", "key"])
wcsv("VOCABULARY.csv", [dict(vocabulary_id=v, vocabulary_name=n, license=l, vocabulary_version=ver) for v, n, l, ver in [("KAI", "K-AI 자체 어휘", "Apache-2.0", "2026.09"), ("SNOMED", "SNOMED CT", "SNOMED International 회원국 라이선스", "확인 필요"), ("LOINC", "LOINC", "무료", "확인 필요"), ("UCUM", "UCUM 단위", "무료", "2.1"), ("MedDRA", "MedDRA", "별도 라이선스", "확인 필요"), ("KCD", "한국표준질병사인분류", "공개", "8차"), ("EDI", "건강보험 EDI 코드", "공개", "확인 필요")]], ["vocabulary_id", "vocabulary_name", "license", "vocabulary_version"])
rel = []
for cs in concept_sets.values():
    for it in cs["items"]: rel.append(dict(concept_id_1=it["concept_id"], concept_id_2=cs["parent_concept_id"], relationship_id="Is a"))
for c in concepts.values():
    if c["vocabulary_id"] != "KAI": rel.append(dict(concept_id_1=c["concept_id"], concept_id_2="", relationship_id="Maps to", target_vocabulary=c["vocabulary_id"], target_code=c["concept_code"]))
wcsv("CONCEPT_RELATIONSHIP.csv", rel, ["concept_id_1", "concept_id_2", "relationship_id", "target_vocabulary", "target_code"])
wcsv("CONCEPT_SET.csv", [dict(concept_set_id=k, concept_set_name_ko=v["kor"], parent_concept_id=v["parent_concept_id"], note=v.get("note") or "") for k, v in concept_sets.items()], ["concept_set_id", "concept_set_name_ko", "parent_concept_id", "note"])
wcsv("CONCEPT_SET_ITEM.csv", [dict(concept_set_id=k, concept_id=it["concept_id"], code=it["code"], label=it["label"]) for k, v in concept_sets.items() for it in v["items"]], ["concept_set_id", "concept_id", "code", "label"])

# ================================================================ 4. 저장 테이블 명세
tables = collections.OrderedDict()
for name, (kor, grain, pk, cols) in OMOP_TABLES.items():
    tables[name] = dict(name=name, layer="OMOP", kor=kor, grain=grain, pk=pk, desc="OMOP CDM 5.4 표준 구조 (concept_id 는 K-AI 어휘)", cohorts=ALL_COHORTS,
                        fields=[dict(name=c, type=t, required=req, desc=d, **({"concept_set": "DISCHARGE_STATUS"} if c == "discharge_to_concept_id" else {})) for c, t, req, d in cols])
def ext_table(F):
    cols = []
    for f in F["fields"].values():
        base = dict(kor=f["kor"], tier=f["tier"], required=bool(f.get("required")), cohorts=sorted(f["cohorts"]))
        if f.get("key"): cols.append(dict(name=f["name"], type="bigint", key=f["key"], ref=f.get("ref"), **base)); continue
        vk = f["value_kind"]
        if vk in ("닫힌 범주", "열린 표준 코드"):
            cols.append(dict(name=f["name"] + "_concept_id", type="bigint", concept_set=f.get("codelist"), vocabulary=f.get("vocabulary"), **base))
            cols.append(dict(name=f["name"] + "_source_value", type="varchar", desc="병원이 고른 코드·라벨 원천값", kor=f["kor"] + " 원천값", tier="권고", required=False, cohorts=sorted(f["cohorts"])))
        elif vk in ("예/아니오", "예/아니오+날짜"): cols.append(dict(name=f["name"], type="integer", concept_set="YN", **base))
        else: cols.append(dict(name=f["name"], type=f["type"], unit=f.get("unit"), pattern=f.get("pattern"), event_date=f.get("event_date"), **base))
    return cols
for fn_ in EXT_FORMS:
    F = forms[fn_]
    tables[fn_] = dict(name=fn_, layer="K-AI 확장", kor=F["kor"], grain=F["grain"], pk=F["pk"], parent=F.get("parent"), desc=F["desc"], cohorts=sorted(F["cohorts"]), fields=ext_table(F))

# ================================================================ 5. 서식 → 레코드 매핑 (필드별 target)
def resolve_target(form, f):
    M = FORM_MAP.get(form)
    if not M: return dict(kind="column", table=form, column=f["name"] + ("_concept_id" if f["value_kind"] in ("닫힌 범주", "열린 표준 코드") else ""))
    n = f["name"]
    for a in M["anchors"]:
        for col, src in a["columns"].items():
            if isinstance(src, str) and re.search(rf"\${n}(:|\||$)", src): return dict(kind="anchor", table=a["table"], column=col, ref=a["ref"])
    ov = (M.get("overrides") or {}).get(n)
    if ov and ov.get("role") == "skip": return dict(kind="skip")
    if ov and ov.get("role") == "episode": return dict(kind="episode", table="EPISODE", concept=ov["concept"], start=ov["start"], parent=ov.get("parent"), object=ov.get("object"))
    if n in M.get("skip", []) or f.get("key"): return dict(kind="skip", note="키·기준 필드")
    table = (ov or {}).get("table") or M["attr_default"]
    if f["value_kind"] == "날짜" and not ov: table = "OBSERVATION"
    val = "value_as_number" if f["value_kind"] == "수치" else ("value_as_concept_id" if f["value_kind"] in ("닫힌 범주", "예/아니오", "예/아니오+날짜") else ("value_as_string" if table == "OBSERVATION" else "value_source_value"))
    link_ref, link_field = M["link"]
    return dict(kind="attr", table=table, concept=f.get("concept_id"), value_column=val, date=M["date_field"], event_ref=link_ref, event_field=link_field, unit=f.get("unit"))
for F in forms.values():
    for f in F["fields"].values(): f["target"] = resolve_target(F["name"], f)

# ================================================================ 6. 예시 투영 (폐암 환자)
def cs_lookup(codelist, code):
    cs = concept_sets.get(codelist)
    if not cs: return None
    for it in cs["items"]:
        if it["code"] == str(code) or it["label"] == str(code): return it["concept_id"]
    return None
BASE_ID = {"CONDITION_OCCURRENCE": 5000, "EPISODE": 7000, "MEASUREMENT": 9000, "OBSERVATION": 8000, "PROCEDURE_OCCURRENCE": 4000, "DRUG_EXPOSURE": 3000, "SPECIMEN": 6000, "OBSERVATION_PERIOD": 1000}
class Proj:
    def __init__(self): self.ids = collections.Counter(); self.records = []; self.refs = {}
    def nid(self, table): self.ids[table] += 1; return BASE_ID.get(table, 0) + self.ids[table]
    def val(self, src, row, form, cohort):
        if not isinstance(src, str): return src
        if src.startswith("#"): return self.refs.get(src[1:].split("(")[0], f"<{src[1:]}>")
        for alt in src.split("|"):
            alt = alt.strip()
            if alt.startswith("@"):
                k = alt[1:].replace("{cohort}", cohort); return registry.get(k, f"<{k}>")
            if alt.startswith("$"):
                fld, _, mode = alt[1:].partition(":"); v = row.get(fld)
                if v in (None, ""): continue
                if mode == "concept":
                    f = forms[form]["fields"].get(fld); c = cs_lookup(f["codelist"], v) if f and f.get("codelist") else None
                    return c if c else f"<concept: {v}>"
                return v
            return alt
        return None
    def add(self, table, cols, ref=None):
        pk = OMOP_TABLES[table][2]; rec = collections.OrderedDict()
        if "+" not in pk: rec[pk] = self.nid(table)
        rec.update(cols); self.records.append((table, rec))
        if ref and "+" not in pk: self.refs[ref] = rec[pk]
        return rec
    def project(self, form, row, cohort):
        M = FORM_MAP[form]; F = forms[form]
        for a in M["anchors"]:
            cols = collections.OrderedDict((c, self.val(s, row, form, cohort)) for c, s in a["columns"].items())
            self.add(a["table"], {k: v for k, v in cols.items() if v not in (None, "")}, a["ref"])
        date = row.get(M["date_field"]) or row.get("index_date")
        for f in F["fields"].values():
            t = f["target"]; v = row.get(f["name"])
            if v in (None, "") or t["kind"] in ("skip", "anchor", "column"): continue
            if t["kind"] == "episode":
                self.add("EPISODE", {k: v2 for k, v2 in dict(person_id=row["person_id"], episode_concept_id=registry[t["concept"]], episode_start_date=self.val(t["start"], row, form, cohort), episode_parent_id=self.val(t["parent"], row, form, cohort) if t.get("parent") else None, episode_object_concept_id=self.val(t["object"], row, form, cohort) if t.get("object") else None, episode_type_concept_id=registry["TYPE_REGISTRY"]).items() if v2 not in (None, "")})
                continue
            tab = t["table"]; ref = next((r for r in t["event_ref"].split("|") if r in self.refs), None)
            rec = collections.OrderedDict(person_id=row["person_id"])
            C = {"MEASUREMENT": ("measurement_concept_id", "measurement_date", "measurement_type_concept_id", "measurement_event_id", "meas_event_field_concept_id"), "OBSERVATION": ("observation_concept_id", "observation_date", "observation_type_concept_id", "observation_event_id", "obs_event_field_concept_id")}[tab]
            rec[C[0]] = t["concept"]; rec[C[1]] = date; rec[C[2]] = registry["TYPE_REGISTRY"]
            if t["value_column"] == "value_as_concept_id":
                c = cs_lookup(f.get("codelist"), v); rec["value_as_concept_id"] = c if c else f"<concept: {v}>"; rec["value_source_value"] = str(v)
            elif t["value_column"] == "value_as_number":
                rec["value_as_number"] = v
                if t.get("unit"): rec["unit_concept_id"] = registry[f"UNIT:{t['unit']}"]
            else: rec[t["value_column"]] = str(v)
            if ref: rec[C[3]] = self.refs[ref]; rec[C[4]] = registry[t["event_field"]]
            self.add(tab, rec)
        for L in M.get("links", []):
            self.records.append((L["table"], collections.OrderedDict((c, s if s.startswith("DRUG_EXPOSURE") else self.val(s, row, form, cohort)) for c, s in L["columns"].items())))

EXAMPLE = [
    ("COHORT_INDEX", dict(person_id=100001, index_date="2022-03-25", diagnosis_registered_date="2022-03-25", diagnosis_name="C34.1", last_followup_date="2023-06-30", last_vital_status="생존")),
    ("STAGING", dict(person_id=100001, staging_date="2022-03-27", initial_clinical_stage="IV", clinical_tnm="cT2aN2M1b", imaging_distant_metastasis_yn=1, distant_metastasis_site="Adrenal gland", distant_metastasis_date="2022-03-25")),
    ("ANTP_THERAPY", dict(person_id=100001, antp_id=1, regimen_or_line_of_therapy=1, regimen="osimertinib", antp_start_date="2022-04-10", antp_end_date="2023-01-15", number_of_cycles=9, treatment_intent_type="고식")),
    ("RESPONSE_ASSESSMENT", dict(person_id=100001, antp_id=1, response_assessment_date="2022-07-12", response_criteria="RECIST 1.1", response_timepoint="3개월", best_overall_response="PR", measurable_disease_yn=1, target_lesions_count=2, spd_mm2=820)),
    ("ADVERSE_EVENT", dict(person_id=100001, antp_id=1, adverse_event_start_date="2022-05-02", adverse_event_concept_id="10037844", adverse_event_source_value="Rash", severity_concept_id="2", causality_concept_id="가능", is_serious=0)),
    ("FAMILY_HISTORY", dict(person_id=100001, record_date="2022-03-25", family_relation="모", family_cancer_type="폐암", family_member_age_at_diagnosis=62)),
    ("SURGERY", dict(person_id=100001, surgery_date="2022-05-20", surgery_name="흉강경 폐엽절제술", surgery_purpose="근치", surgery_approach="흉강경", surgery_site="우중엽", ebl_ml=150)),
    ("RADIOTHERAPY", dict(person_id=100001, rt_start_date="2022-08-01", rt_name="흉부 방사선치료", rt_purpose="고식", rt_site="흉부", rt_dose_per_fraction_gy=3.0, rt_fractions=10, rt_total_dose_gy=30.0)),
]
P = Proj()
ex = ["# 서식 → 레코드 예시 (폐암 환자 100001)", "", f"`crf_forms_v2.yaml` 의 매핑 규칙을 `build_spec_v2.py` 가 실제로 적용해 만든 결과입니다. concept_id 는 `vocab/CONCEPT.csv` 의 K-AI 어휘이며, `<concept: …>` 는 아직 값 집합에 없는 값(2.1 에서 채움), `<…>` 는 다른 서식이 만든 레코드 참조입니다.", ""]
for form, row in EXAMPLE:
    n0 = len(P.records); P.project(form, row, "LUNG_CANCER")
    ex += [f"## {form} · {FORM_MAP[form]['kor']}", "", "입력 (서식 한 행):", "", "```", ", ".join(f"{k}={v}" for k, v in row.items()), "```", "", "생성 레코드:", ""]
    for table, rec in P.records[n0:]: ex.append(f"- **{table}** " + ", ".join(f"{k}={v}" for k, v in rec.items() if v not in (None, "")))
    ex.append("")
ex += ["## 규칙 요약", "", "| 서식 | 기준 레코드 | 속성 레코드 | 속성의 연결 대상 |", "|---|---|---|---|"]
for form in OMOP_FORMS:
    M = FORM_MAP[form]; anc = ", ".join(dict.fromkeys(a["table"] for a in M["anchors"])) or "(없음 · 속성만)"
    ex.append(f"| {form} | {anc} | {M['attr_default']} (필드 concept + 값) | {M['link'][0]} → {M['link'][1]} |")
ex += ["", "속성 레코드 규칙: 서식의 나머지 필드마다 레코드 하나. `*_concept_id` = 필드 concept(CONCEPT.csv 의 Attribute), 값은 수치 → `value_as_number`(+unit), 범주 → `value_as_concept_id`(+`value_source_value`), 텍스트 → `value_as_string`. 날짜는 서식의 기준일, `*_event_id` 는 기준 레코드 id, `*_event_field_concept_id` 는 그 테이블의 Field concept."]
open(os.path.join(OUT, "crf_to_omop_examples.md"), "w", encoding="utf-8").write("\n".join(ex) + "\n")

# ================================================================ 7. YAML·엑셀·보고서
class D(yaml.SafeDumper): pass
D.add_representer(collections.OrderedDict, lambda d, data: d.represent_dict(data.items()))
def dump(obj, fn):
    with open(os.path.join(OUT, fn), "w", encoding="utf-8") as fh:
        yaml.dump(json.loads(json.dumps(obj, ensure_ascii=False, default=lambda o: sorted(o) if isinstance(o, set) else str(o))), fh, Dumper=D, allow_unicode=True, sort_keys=False, width=140)
def clean(f):
    o = collections.OrderedDict()
    for k in ("name", "kor", "type", "value_kind", "tier", "required", "key", "ref", "event_date", "codelist", "vocabulary", "unit", "pattern", "standard", "concept_id", "derived_score", "desc", "target"):
        v = f.get(k)
        if v not in (None, False, "", []): o[k] = v
    o["cohorts"] = sorted(f["cohorts"])
    if f.get("also"): o["also"] = f["also"]
    if f.get("v1_hint"): o["v1_hint"] = f["v1_hint"]
    return dict(o)
spec = collections.OrderedDict(meta=dict(name="K-AI 질환별 코호트 데이터 모델", version="2.0.0-draft", source=os.path.basename(XLSX), generated_by="spec/build_spec_v2.py", generated_at=TODAY,
    concept_id="K-AI 어휘 · 64비트 정수 · 10,000,000,000 부터 · vocab/CONCEPT.csv",
    principles=["저장 테이블은 OMOP CDM 5.4 + Oncology(EPISODE, EPISODE_EVENT)의 이름·컬럼을 그대로 쓴다", "concept_id 는 OMOP 이 아니라 K-AI 어휘이며, 표준 코드(SNOMED·LOINC·MedDRA)는 CONCEPT.concept_code 에 보관한다",
                "병원은 CRF 서식을 채우고 앱이 서식→레코드 규칙으로 저장 테이블을 만든다", "OMOP 에 자리가 없는 병리 세부·질환 항목은 K-AI 확장 테이블(OMOP 관례: *_concept_id + *_source_value)", "한 테이블 = 행 단위 하나"]),
    layers=dict(omop=[t for t, T in tables.items() if T["layer"] == "OMOP"], extension=[t for t, T in tables.items() if T["layer"] != "OMOP"]),
    tables=tables)
dump(spec, "kai_cdm_spec_v2.yaml")
dump(dict(forms=collections.OrderedDict((n, dict(name=n, kor=F["kor"], grain=F["grain"], kind="OMOP 투영" if n in FORM_MAP else "확장 테이블 직접 저장",
     storage=list(dict.fromkeys([a["table"] for a in FORM_MAP[n]["anchors"]] + [FORM_MAP[n]["attr_default"]])) if n in FORM_MAP else [n], cohorts=sorted(F["cohorts"]), fields=[clean(f) for f in F["fields"].values()])) for n, F in forms.items())), "crf_forms_v2.yaml")
dump(dict(cohorts=collections.OrderedDict((c, dict(kor=COHORT_KOR[c], omop_tables=list(OMOP_TABLES), extension_tables=[t for t in EXT_FORMS if c in forms[t]["cohorts"]], forms=[n for n, F in forms.items() if c in F["cohorts"]])) for c in ALL_COHORTS)), "cohort_profiles_v2.yaml")
dump(dict(derived=derived, derived_scores=scores), "derived_v2.yaml")
merged = {h for F in forms.values() for f in F["fields"].values() for h in f["v1_hint"]}
legacy = collections.defaultdict(list)
for r in V1FIELDS:
    key = f"{r['테이블']}.{r['필드']}"
    if r["처리"] in ("이동", "이관→코어") and key not in merged: legacy[r["목적지 테이블"]].append(dict(v1=key, kor=r.get("한글명") or "", desc=(r.get("설명") or "")[:120], type=r.get("타입"), action=r["처리"]))
dump(dict(legacy=dict(legacy)), "legacy_fields_v2.yaml")

xw = Workbook(); H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="003B73"); WRAP = Alignment(wrap_text=True, vertical="top")
def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]: c.font = H; c.fill = HF; c.alignment = WRAP
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
def target_text(t):
    if t["kind"] in ("anchor", "column"): return f"{t['table']}.{t['column']}"
    if t["kind"] == "attr": return f"{t['table']} 속성 ({t['value_column']}, {t['event_ref']}→{t['event_field']})"
    if t["kind"] == "episode": return f"EPISODE({t['concept']})"
    return "키/기준"
ws = xw.active; ws.title = "저장테이블"
sheet(ws, ["테이블", "층", "한글", "행 단위", "고유키", "필드 수", "설명"], [[n, T["layer"], T["kor"], T["grain"], T["pk"], len(T["fields"]), T["desc"]] for n, T in tables.items()], [26, 10, 18, 24, 26, 7, 70])
sheet(xw.create_sheet("저장컬럼"), ["테이블", "컬럼", "타입", "필수", "한글", "값 집합", "참조", "설명"], [[n, f["name"], f["type"], "Y" if f.get("required") else "", f.get("kor") or "", f.get("concept_set") or "", f.get("ref") or "", f.get("desc") or ""] for n, T in tables.items() for f in T["fields"]], [26, 34, 9, 6, 28, 22, 34, 50])
sheet(xw.create_sheet("서식"), ["서식", "종류", "필드", "한글", "타입", "값 종류", "등급", "코드표(값 집합)", "단위", "저장 위치", "적용 코호트", "같은 뜻 항목"],
      [[n, "OMOP 투영" if n in FORM_MAP else "확장 저장", f["name"], f["kor"], f["type"], f["value_kind"], f["tier"], f.get("codelist") or "", f.get("unit") or "", target_text(f["target"]), ", ".join(COHORT_KOR[c] for c in sorted(f["cohorts"])), "; ".join(f.get("also") or [])] for n, F in forms.items() for f in F["fields"].values()], [24, 10, 30, 30, 8, 12, 6, 22, 8, 46, 26, 30])
sheet(xw.create_sheet("CONCEPT"), ["concept_id", "한글", "영문", "domain", "vocabulary", "class", "code", "standard", "표준 매핑 후보", "key"], [[c["concept_id"], c["concept_name_ko"], c["concept_name_en"], c["domain_id"], c["vocabulary_id"], c["concept_class_id"], c["concept_code"], c["standard_concept"], c["standard_hint"], c["key"]] for c in concepts.values()], [14, 30, 30, 12, 10, 12, 22, 8, 30, 40])
sheet(xw.create_sheet("CONCEPT_SET"), ["값 집합", "한글", "concept_id", "코드", "라벨", "비고"], [[k, v["kor"], it["concept_id"], it["code"], it["label"], v.get("note") or ""] for k, v in concept_sets.items() for it in (v["items"] or [dict(concept_id="", code="", label="(비어 있음)")])], [26, 24, 14, 10, 30, 36])
sheet(xw.create_sheet("파생변수"), ["변수", "한글", "코호트", "원천 테이블", "값 종류", "등급", "concept set", "표준", "메모"], [[d["name"], d["kor"], COHORT_KOR[d["cohort"]], d["source_table"], d["value_kind"], d["tier"], d.get("concept_set") or "", d.get("standard") or "", d.get("note") or ""] for d in derived], [28, 32, 10, 22, 14, 6, 20, 26, 40])
sheet(xw.create_sheet("v1 미승격 필드"), ["목적지", "v1 필드", "한글", "타입", "처리", "설명"], [[d, e["v1"], e["kor"], e["type"], e["action"], e["desc"]] for d, es in legacy.items() for e in es], [22, 40, 28, 8, 12, 60])
xw.save(os.path.join(OUT, "dictionary_v2.xlsx"))

nfields = sum(len(F["fields"]) for F in forms.values()); tiers = collections.Counter(f["tier"] for F in forms.values() for f in F["fields"].values())
L = [f"# 사전 v2 변환 보고서 ({TODAY})", "", f"입력: `{os.path.relpath(XLSX, ROOT)}` · 출력: `spec/v2/`", "", "## 규모", "", "| 구분 | 수 |", "|---|---|",
     f"| 저장 테이블 | {len(tables)} (OMOP {len(OMOP_TABLES)}, K-AI 확장 {len(tables) - len(OMOP_TABLES)}) |", f"| CRF 서식 | {len(forms)} (OMOP 투영 {len(OMOP_FORMS)}, 확장 저장 {len(EXT_FORMS)}) |",
     f"| 서식 필드 | {nfields} (필수 {tiers['필수']}, 권고 {tiers['권고']}, 보류 {tiers['보류']}) |", f"| K-AI concept | {len(concepts)} (SNOMED 코드 보유 {sum(1 for c in concepts.values() if c['vocabulary_id'] == 'SNOMED')}) |",
     f"| 값 집합 | {len(concept_sets)} (비어 있음 {sum(1 for v in concept_sets.values() if not v['items'])}) |", f"| 코어 파생 변수 | {len(derived)} · 파생 점수 {len(scores)} |", f"| v1 미승격 필드 | {sum(len(v) for v in legacy.values())} |", "",
     "## 서식과 저장 위치", "", "| 서식 | 행 단위 | 필드 | 저장 |", "|---|---|---|---|"]
for n, F in forms.items():
    st = (", ".join(dict.fromkeys([a["table"] for a in FORM_MAP[n]["anchors"]] + [FORM_MAP[n]["attr_default"]])) if n in FORM_MAP else f"{n} (확장 테이블)")
    L.append(f"| {n} | {F['grain']} | {len(F['fields'])} | {st} |")
L += ["", "## 비어 있는 값 집합 (2.1 에서 채움)", "", ", ".join(k for k, v in concept_sets.items() if not v["items"]) or "없음", "", "## v1 미승격 필드", ""] + [f"- {d}: {len(es)}개" for d, es in sorted(legacy.items(), key=lambda x: -len(x[1]))] + ["", "## 경고", ""] + [f"- {w}" for w in warnings]
open(os.path.join(OUT, "spec_v2_report.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")
print("\n".join(L[:15])); print(f"... 경고 {len(warnings)}건 → {OUT}")
